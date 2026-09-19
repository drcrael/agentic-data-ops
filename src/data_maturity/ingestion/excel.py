"""Inspect OOXML workbooks without executing macros, links, or formulas."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from data_maturity.config import Config
from data_maturity.ingestion.base import IngestionError, collect, make_dataset, source_metadata
from data_maturity.ingestion.table_detection import DetectedTable, detect_tables
from data_maturity.models.core import Dataset, DatasetStructure, SourceMetadata


class ExcelAdapter:
    def read(self, path: Path, config: Config) -> tuple[SourceMetadata, list[Dataset]]:
        source = source_metadata(path, config)
        try:
            with ZipFile(path) as archive:
                if (
                    sum(item.file_size for item in archive.infolist())
                    > config.profiling.max_uncompressed_bytes
                ):
                    raise IngestionError("Workbook exceeds max_uncompressed_bytes")
                cells = 0
                for name in archive.namelist():
                    if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                        # Count serialized cells before building an object graph.
                        data = archive.read(name)
                        cells += data.count(b"<c ") + data.count(b"<c>")
                if cells > config.profiling.max_excel_cells:
                    raise IngestionError("Workbook exceeds max_excel_cells")
            workbook = load_workbook(path, data_only=False, keep_links=False, keep_vba=False)
        except (BadZipFile, KeyError, OSError, ValueError) as exc:
            if isinstance(exc, IngestionError):
                raise
            raise IngestionError(f"Unable to read workbook: {type(exc).__name__}") from exc
        datasets: list[Dataset] = []
        metadata: list[dict[str, Any]] = []
        try:
            for sheet in workbook.worksheets:
                if (
                    sheet.max_column > config.profiling.max_columns
                    or sheet.max_row * sheet.max_column > config.profiling.max_excel_cells
                ):
                    raise IngestionError("Worksheet dimensions exceed configured inspection limits")
                regions = detect_tables(sheet)
                formulas = [
                    cell.coordinate for row in sheet for cell in row if cell.data_type == "f"
                ]
                metadata.append(
                    {
                        "name": sheet.title,
                        "state": sheet.sheet_state,
                        "dimensions": sheet.calculate_dimension(),
                        "tables": {n: sheet.tables[n].ref for n in sheet.tables},
                        "merged_cells": [str(r) for r in sheet.merged_cells.ranges],
                        "hidden_rows": [i for i, d in sheet.row_dimensions.items() if d.hidden],
                        "hidden_columns": [
                            i for i, d in sheet.column_dimensions.items() if d.hidden
                        ],
                        "formula_count": len(formulas),
                        "formula_locations": formulas[:100],
                        "data_validations": [
                            {"range": str(v.sqref), "type": v.type}
                            for v in sheet.data_validations.dataValidation
                        ],
                        "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
                        "filter": sheet.auto_filter.ref,
                        "discovered_regions": [r.model_dump() for r in regions],
                    }
                )
                for region in regions:
                    headers = [
                        sheet.cell(region.min_row, c).value
                        for c in range(region.min_col, region.max_col + 1)
                    ]
                    formula_cells: list[str] = []

                    def records(
                        region: DetectedTable = region,
                        sheet: Worksheet = sheet,
                        formula_cells: list[str] = formula_cells,
                    ) -> Iterator[tuple[int, list[Any]]]:
                        for number in range(region.min_row + 1, region.max_row + 1):
                            values: list[Any] = []
                            for col in range(region.min_col, region.max_col + 1):
                                cell = sheet.cell(number, col)
                                if cell.data_type == "f":
                                    formula_cells.append(cell.coordinate)
                                    values.append(
                                        None
                                    )  # Never treat formula expressions as computed values.
                                elif cell.data_type == "e":
                                    values.append(None)
                                else:
                                    values.append(cell.value)
                            yield number, values

                    rows, numbers, total, scanned = collect(records(), config)
                    sampled = total > len(rows)
                    structure = DatasetStructure(
                        worksheet=sheet.title,
                        table=region.name,
                        cell_range=region.cell_range,
                        header_row=region.min_row,
                        confidence=region.confidence,
                        detection_method=region.method,
                        alternatives=region.alternatives,
                        source_row_count=total,
                        rows_scanned=scanned,
                        rows_retained=len(rows),
                        sampled=sampled,
                        truncated=total > scanned,
                        sampling_method="reservoir over scanned prefix" if sampled else "full",
                        seed=config.runtime.random_seed,
                        formula_count=len(formula_cells),
                        formula_locations=formula_cells[:100],
                        warnings=[
                            "Formula results are unavailable; formulas are counted, not evaluated"
                        ]
                        if formula_cells
                        else [],
                    )
                    datasets.append(make_dataset(source, structure, headers, rows, numbers))
            source.inspection = {
                "worksheets": metadata,
                "named_ranges": [
                    {"name": n, "reference": v.attr_text} for n, v in workbook.defined_names.items()
                ],
                "macros_executed": False,
                "formulas_evaluated": False,
            }
            for dataset in datasets:
                dataset.source = source
        finally:
            workbook.close()
        if not datasets:
            raise IngestionError("No tabular datasets discovered")
        return source, datasets
