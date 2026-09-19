"""Conservative region discovery with explicit ranges and blank separators."""

from __future__ import annotations

from typing import Any

from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import Field

from data_maturity.models.core import Model


class DetectedTable(Model):
    name: str
    min_row: int
    max_row: int
    min_col: int
    max_col: int
    confidence: float = Field(ge=0, le=1)
    method: str
    alternatives: list[str] = Field(default_factory=list)

    @property
    def cell_range(self) -> str:
        return f"{get_column_letter(self.min_col)}{self.min_row}:{get_column_letter(self.max_col)}{self.max_row}"


def groups(indices: list[int]) -> list[list[int]]:
    result: list[list[int]] = []
    for index in indices:
        if not result or index != result[-1][-1] + 1:
            result.append([index])
        else:
            result[-1].append(index)
    return result


def detect_tables(sheet: Worksheet) -> list[DetectedTable]:
    regions: list[DetectedTable] = []
    for name in sheet.tables:
        table = sheet.tables[name]
        left, top, right, bottom = range_boundaries(table.ref)
        if None in (left, top, right, bottom):
            raise ValueError("Excel table must have a bounded rectangular range")
        assert left is not None and top is not None and right is not None and bottom is not None
        regions.append(
            DetectedTable(
                name=name,
                min_row=top,
                max_row=bottom,
                min_col=left,
                max_col=right,
                confidence=1,
                method="excel_table",
            )
        )

    def covered(row: int, col: int) -> bool:
        return any(t.min_row <= row <= t.max_row and t.min_col <= col <= t.max_col for t in regions)

    # Only occupied, uncovered cells drive inferred boundaries. Formatting-only cells do not.
    occupied: dict[int, dict[int, Any]] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None and not covered(cell.row, cell.column):
                occupied.setdefault(cell.row, {})[cell.column] = cell.value
    for band in groups(sorted(occupied)):
        columns = sorted({c for r in band for c in occupied[r]})
        for cols in groups(columns):
            left, right = cols[0], cols[-1]
            active = [r for r in band if any(left <= c <= right for c in occupied[r])]
            if not active:
                continue
            width = right - left + 1
            # Strip titles/notes, not sparse interior records. Ambiguous trimming is disclosed.
            trimmed = False
            while (
                len(active) > 1
                and width > 1
                and sum(left <= c <= right for c in occupied[active[0]]) < 2
            ):
                active.pop(0)
                trimmed = True
            while (
                len(active) > 1
                and width > 1
                and sum(left <= c <= right for c in occupied[active[-1]]) < 2
            ):
                val = next(iter(occupied[active[-1]].values()))
                if isinstance(val, str) and (
                    len(val) > 35 or val.lower().startswith(("note", "source:", "total"))
                ):
                    active.pop()
                    trimmed = True
                else:
                    break
            if len(active) < 2:
                continue
            top, bottom = active[0], active[-1]
            headers = [sheet.cell(top, c).value for c in range(left, right + 1)]
            repeated = [
                r
                for r in active[1:]
                if [sheet.cell(r, c).value for c in range(left, right + 1)] == headers
            ]
            starts = [top, *repeated]
            for index, start in enumerate(starts):
                end = starts[index + 1] - 1 if index + 1 < len(starts) else bottom
                if end <= start:
                    continue
                string_headers = all(isinstance(v, str) and v.strip() for v in headers)
                confidence = 0.9 if string_headers and len(set(headers)) == width else 0.6
                alternatives = ["Boundary/header inferred; confirm against source documentation"]
                if trimmed:
                    alternatives.append("Single-cell title or note excluded heuristically")
                regions.append(
                    DetectedTable(
                        name=f"region_{start}_{left}",
                        min_row=start,
                        max_row=end,
                        min_col=left,
                        max_col=right,
                        confidence=confidence,
                        method="blank_boundaries_and_header",
                        alternatives=alternatives,
                    )
                )
    return sorted(regions, key=lambda t: (t.min_row, t.min_col))
